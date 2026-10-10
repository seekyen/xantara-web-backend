"""Downloadable stock movement report (Excel / PDF) for one product at one branch.

Rows come from StockMovement, the complete ledger (receiving, sales, returns,
write-offs, transfers, adjustments). Batch numbers are resolved through the
shared movement reference when the movement touched a batch.
"""
from datetime import datetime, time
from io import BytesIO
from django.utils import timezone
from rest_framework import serializers
from .models import StockMovement, InventoryBatchMovement

TYPE_LABELS = dict(StockMovement.MOVEMENT_TYPES)
LOCATION_LABELS = dict(StockMovement.LOCATION_CHOICES)
HEADERS = ['Date / time (PH)', 'Type', 'Batch', 'Location', 'In', 'Out', 'Before', 'After', 'Reference', 'Reason / remarks', 'By']
MAX_ROWS = 20000


def _number(value):
    return int(value) if float(value).is_integer() else round(float(value), 4)


def parse_range(params):
    """Inclusive Philippine-time calendar dates -> aware datetimes, or None."""
    tz = timezone.get_current_timezone()
    bounds = []
    for key, end in (('date_from', False), ('date_to', True)):
        raw = params.get(key)
        if not raw:
            bounds.append(None)
            continue
        try:
            day = datetime.strptime(raw, '%Y-%m-%d').date()
        except ValueError:
            raise serializers.ValidationError({key: 'Use the format YYYY-MM-DD.'})
        bounds.append(timezone.make_aware(datetime.combine(day, time.max if end else time.min), tz))
    if bounds[0] and bounds[1] and bounds[0] > bounds[1]:
        raise serializers.ValidationError({'date_to': 'End date cannot precede the start date.'})
    return bounds


def movement_report(product, branch, start, end):
    from apps.accounts.models import Staff
    rows = StockMovement.objects.filter(itemcode=product.itemcode, branch_code=branch)
    if start:
        rows = rows.filter(created_at__gte=start)
    if end:
        rows = rows.filter(created_at__lte=end)
    rows = list(rows.order_by('created_at', 'pk')[:MAX_ROWS + 1])
    truncated = len(rows) > MAX_ROWS
    rows = rows[:MAX_ROWS]
    batch_by_ref = dict(InventoryBatchMovement.objects.filter(batch__product=product, branch_code=branch)
                        .values_list('reference', 'batch__number'))
    staff_ids = {int(r.created_by) for r in rows if r.created_by.isdigit()}
    staff = dict(Staff.objects.filter(pk__in=staff_ids).values_list('pk', 'name'))
    lines, totals = [], {}
    total_in = total_out = 0
    for r in rows:
        qty = _number(r.qty)
        total_in += qty if qty > 0 else 0
        total_out += -qty if qty < 0 else 0
        entry = totals.setdefault(r.movement_type, {'count': 0, 'qty': 0})
        entry['count'] += 1
        entry['qty'] += qty
        by = staff.get(int(r.created_by), r.created_by) if r.created_by.isdigit() else r.created_by
        lines.append([timezone.localtime(r.created_at).strftime('%Y-%m-%d %H:%M'), TYPE_LABELS.get(r.movement_type, r.movement_type),
                      batch_by_ref.get(r.ref_no, ''), LOCATION_LABELS.get(r.location, r.location),
                      qty if qty > 0 else '', -qty if qty < 0 else '', _number(r.qty_before), _number(r.qty_after),
                      r.ref_no, r.remarks, by])
    period = ' to '.join(v.astimezone(timezone.get_current_timezone()).strftime('%Y-%m-%d') for v in (start, end) if v)
    if start and not end:
        period = 'From ' + period
    elif end and not start:
        period = 'Up to ' + period
    return {
        'title': f'Stock movement report — {product.desclong or product.descshort or product.itemcode}',
        'meta': [('Item code', product.itemcode), ('Branch', branch), ('Period', period or 'All dates'),
                 ('Generated', timezone.localtime().strftime('%Y-%m-%d %H:%M')),
                 ('Movements', f'{len(lines):,}' + (f' (first {MAX_ROWS:,} shown; narrow the dates)' if truncated else ''))],
        'summary': [[TYPE_LABELS.get(k, k), v['count'], v['qty']] for k, v in sorted(totals.items())],
        'total_in': total_in, 'total_out': total_out, 'net': total_in - total_out, 'rows': lines,
    }


def build_xlsx(report):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = Workbook()
    sheet = book.active
    sheet.title = 'Movements'
    sheet['A1'] = report['title']
    sheet['A1'].font = Font(bold=True, size=14)
    for i, (label, value) in enumerate(report['meta'], start=2):
        sheet.cell(i, 1, label).font = Font(bold=True)
        sheet.cell(i, 2, value)
    header_row = len(report['meta']) + 3
    for col, name in enumerate(HEADERS, start=1):
        cell = sheet.cell(header_row, col, name)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1D4ED8')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r, line in enumerate(report['rows'], start=header_row + 1):
        for col, value in enumerate(line, start=1):
            # Text is written as a string; a reason starting with "=" must never run as a formula.
            cell = sheet.cell(r, col, value)
            if isinstance(value, str):
                cell.data_type = 's'
    last = header_row + len(report['rows'])
    total_row = last + 1
    sheet.cell(total_row, 4, 'Totals').font = Font(bold=True)
    for col in (5, 6):
        letter = get_column_letter(col)
        cell = sheet.cell(total_row, col, f'=SUM({letter}{header_row + 1}:{letter}{max(last, header_row + 1)})')
        cell.font = Font(bold=True)
    sheet.cell(total_row + 1, 4, 'Net change').font = Font(bold=True)
    sheet.cell(total_row + 1, 5, f'=E{total_row}-F{total_row}').font = Font(bold=True)
    sheet.freeze_panes = sheet.cell(header_row + 1, 1)
    sheet.auto_filter.ref = f'A{header_row}:{get_column_letter(len(HEADERS))}{max(last, header_row + 1)}'
    for col, width in enumerate([18, 18, 18, 13, 9, 9, 9, 9, 26, 40, 20], start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    summary = book.create_sheet('Summary')
    summary['A1'] = report['title']
    summary['A1'].font = Font(bold=True, size=14)
    for col, name in enumerate(['Movement type', 'Count', 'Net quantity'], start=1):
        summary.cell(3, col, name).font = Font(bold=True)
    for r, row in enumerate(report['summary'], start=4):
        for col, value in enumerate(row, start=1):
            summary.cell(r, col, value)
    base = len(report['summary']) + 5
    for i, (label, value) in enumerate([('Total in', report['total_in']), ('Total out', report['total_out']), ('Net change', report['net'])]):
        summary.cell(base + i, 1, label).font = Font(bold=True)
        summary.cell(base + i, 3, value).font = Font(bold=True)
    summary.column_dimensions['A'].width = 26
    summary.column_dimensions['B'].width = 10
    summary.column_dimensions['C'].width = 16
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def build_pdf(report):
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    styles = getSampleStyleSheet()
    cell = ParagraphStyle('cell', parent=styles['BodyText'], fontSize=7, leading=8.5)
    head = ParagraphStyle('head', parent=cell, textColor=colors.white, fontName='Helvetica-Bold')
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                            title=report['title'])
    story = [Paragraph(escape(report['title']), styles['Title'])]
    story.append(Paragraph('<br/>'.join(f'<b>{escape(k)}:</b> {escape(str(v))}' for k, v in report['meta']), styles['BodyText']))
    story.append(Spacer(1, 4 * mm))
    summary = [['Movement type', 'Count', 'Net quantity']] + [[a, b, c] for a, b, c in report['summary']]
    summary.append(['Total in / out / net', f"{report['total_in']} / {report['total_out']}", report['net']])
    table = Table(summary, colWidths=[60 * mm, 40 * mm, 30 * mm], hAlign='LEFT')
    table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1D4ED8')), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                               ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
                               ('FONTSIZE', (0, 0), (-1, -1), 8), ('GRID', (0, 0), (-1, -1), .25, colors.grey), ('ALIGN', (1, 0), (-1, -1), 'RIGHT')]))
    story += [table, Spacer(1, 5 * mm)]
    data = [[Paragraph(h, head) for h in HEADERS]]
    data += [[Paragraph(escape(str(v)), cell) for v in line] for line in report['rows']]
    if not report['rows']:
        data.append([Paragraph('No stock movements in this period.', cell)] + [''] * (len(HEADERS) - 1))
    widths = [26, 24, 22, 18, 12, 12, 13, 13, 36, 62, 27]
    scale = (277 * mm) / (sum(widths) * mm)
    grid = Table(data, colWidths=[w * mm * scale for w in widths], repeatRows=1)
    grid.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1D4ED8')), ('GRID', (0, 0), (-1, -1), .25, colors.lightgrey),
                              ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F4F6')])]))
    story.append(grid)

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont('Helvetica', 7)
        canvas.drawRightString(landscape(A4)[0] - 10 * mm, 6 * mm, f'Page {document.page}')
        canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()
