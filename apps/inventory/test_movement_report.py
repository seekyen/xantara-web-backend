from io import BytesIO
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient
from apps.accounts.models import Staff
from apps.inventory.batches import receive_stock, write_off
from apps.inventory.models import Branch, InventoryBatchBalance, Product
from apps.inventory.workspace import save_document


class MovementReportTests(TestCase):
    def setUp(self):
        Branch.objects.get_or_create(code='MAIN', defaults={'name': 'Main'})
        self.product = Product.objects.get(pk=save_document({'reason': 'New product', 'data': {'name': 'Milk', 'quantity': 0}}, 'admin')['id'])
        self.user = Staff.objects.create_user('report@example.test', role='manager', name='Rita Manager')
        receive_stock(self.product.pk, {'number': 'LOT-1', 'quantity': 12, 'received_date': str(timezone.localdate()), 'unit_cost': '35.00',
                                        'reason': '=HYPERLINK("http://x")'}, str(self.user.pk))
        balance = InventoryBatchBalance.objects.get()
        write_off(self.product.pk, {'branch': 'MAIN', 'balance_id': balance.pk, 'quantity': 1, 'expected_quantity': 12, 'reason': 'Broken'}, str(self.user.pk))
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.url = f'/api/v1/inventory-items/{self.product.pk}/movement-report/'

    def test_excel_contains_movements_totals_and_no_formula_injection(self):
        response = self.client.get(self.url, {'branch': 'MAIN', 'file': 'xlsx'})
        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment; filename="stock-movements-', response['Content-Disposition'])
        book = load_workbook(BytesIO(response.content))
        rows = list(book['Movements'].iter_rows(values_only=True))
        header = next(i for i, row in enumerate(rows) if row[0] == 'Date / time (PH)')
        data = rows[header + 1:header + 3]
        self.assertEqual([r[1] for r in data], ['Stock receiving', 'Write-off / damage'])
        self.assertEqual([r[2] for r in data], ['LOT-1', 'LOT-1'])
        self.assertEqual((data[0][4], data[1][5]), (12, 1))
        self.assertEqual(data[0][9], '=HYPERLINK("http://x")')
        self.assertEqual(book['Movements'].cell(header + 4, 5).value.startswith('=SUM('), True)
        self.assertEqual(data[0][10], 'Rita Manager')

    def test_pdf_and_filters_and_errors(self):
        pdf = self.client.get(self.url, {'file': 'pdf'})
        self.assertEqual((pdf.status_code, pdf['Content-Type']), (200, 'application/pdf'))
        self.assertTrue(pdf.content.startswith(b'%PDF'))
        empty = self.client.get(self.url, {'file': 'pdf', 'date_from': '2001-01-01', 'date_to': '2001-01-31'})
        self.assertEqual(empty.status_code, 200)
        self.assertEqual(self.client.get(self.url, {'file': 'csv'}).status_code, 400)
        self.assertEqual(self.client.get(self.url, {'date_from': 'bad'}).status_code, 400)
        self.assertEqual(self.client.get(self.url, {'date_from': '2026-02-02', 'date_to': '2026-01-01'}).status_code, 400)
        self.assertEqual(self.client.get(self.url, {'branch': 'NOPE'}).status_code, 404)

    def test_cashier_is_refused(self):
        self.client.force_authenticate(Staff.objects.create_user('cash@example.test', role='cashier', name='C'))
        self.assertEqual(self.client.get(self.url).status_code, 403)
